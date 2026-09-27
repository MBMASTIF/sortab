"""End-to-end HTTP tests for the Инструменты endpoints (Unpivot, Compare) —
see backend/main.py's "Инструменты" section and backend/tools_store.py.
Deliberately does NOT touch Redis or Postgres: these tools are stateless
and anonymous by design, so the only override needed is a fresh
ToolFileStore per test (no shared state leaking between tests, same
"override the dependency" pattern test_api.py uses for Redis).
"""

from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.main import app, get_tools_store
from backend.tools_store import ToolFileStore


@pytest.fixture
def client():
    # One shared instance for the whole test — get_tools_store is called
    # once per request, so a lambda that builds a NEW ToolFileStore() every
    # call would silently lose every token between the preview request and
    # the commit request (caught by running this suite for real, not
    # assumed).
    store = ToolFileStore()
    app.dependency_overrides[get_tools_store] = lambda: store
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def _xlsx_bytes(rows: list[list]) -> bytes:
    import openpyxl as _openpyxl

    wb = _openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------- Unpivot ----------------------------


def _unpivot_source_bytes() -> bytes:
    return _xlsx_bytes(
        [
            ["Товар", "Январь", "Февраль", "Март"],
            ["Носки", 1000, 1100, 1200],
            ["Ботинки", 2000, 2200, 2300],
        ]
    )


def test_unpivot_preview_returns_token_and_grid(client):
    resp = client.post(
        "/api/tools/preview",
        files={"file": ("Сводная.xlsx", _unpivot_source_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["token"]
    assert body["row_count"] == 3
    assert body["preview_rows"][0] == ["Товар", "Январь", "Февраль", "Март"]


def test_unpivot_commit_downloads_long_format_xlsx(client):
    preview = client.post(
        "/api/tools/preview",
        files={"file": ("Сводная.xlsx", _unpivot_source_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()

    resp = client.post(
        "/api/tools/unpivot",
        json={
            "token": preview["token"],
            "header_row_index": 0,
            "id_columns": [0],
            "value_columns": [1, 2, 3],
        },
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = openpyxl.load_workbook(BytesIO(resp.content))
    ws = wb["Развёрнуто"]
    header = [c.value for c in ws[1]]
    assert header == ["Товар", "Категория", "Значение"]

    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == 6  # 2 products x 3 months

    total = sum(Decimal(str(r[2])) for r in data_rows)
    expected = Decimal(1000 + 1100 + 1200 + 2000 + 2200 + 2300)
    assert total == expected


def test_unpivot_rejects_fewer_than_two_value_columns(client):
    preview = client.post(
        "/api/tools/preview",
        files={"file": ("Сводная.xlsx", _unpivot_source_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()

    resp = client.post(
        "/api/tools/unpivot",
        json={"token": preview["token"], "header_row_index": 0, "id_columns": [0], "value_columns": [1]},
    )
    assert resp.status_code == 400


def test_unpivot_unknown_token_is_404(client):
    resp = client.post(
        "/api/tools/unpivot",
        json={"token": "does-not-exist", "header_row_index": 0, "id_columns": [0], "value_columns": [1, 2]},
    )
    assert resp.status_code == 404


# ---------------------------- Compare ----------------------------


def _compare_file_a_bytes() -> bytes:
    return _xlsx_bytes(
        [
            ["Контрагент", "Сумма"],
            ["ООО Ромашка", 10000],
            ["ИП Иванов", 5000],
            ["ЗАО Вектор", 777.77],
        ]
    )


def _compare_file_b_bytes() -> bytes:
    return _xlsx_bytes(
        [
            ["Контрагент", "Сумма"],
            ["ООО Ромашка", 10000],
            ["ИП Иванов", 5500],
            ["ООО Заря", 300],
        ]
    )


def _upload_for_compare(client, content: bytes, name: str) -> dict:
    resp = client.post(
        "/api/tools/preview",
        files={"file": (name, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_compare_returns_all_four_statuses(client):
    a = _upload_for_compare(client, _compare_file_a_bytes(), "A.xlsx")
    b = _upload_for_compare(client, _compare_file_b_bytes(), "B.xlsx")

    resp = client.post(
        "/api/tools/compare",
        json={
            "token_a": a["token"],
            "header_row_index_a": 0,
            "entity_column_a": 0,
            "metric_column_a": 1,
            "token_b": b["token"],
            "header_row_index_b": 0,
            "entity_column_b": 0,
            "metric_column_b": 1,
        },
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()

    by_entity = {r["entity"]: r for r in body["rows"]}
    assert by_entity["ООО Ромашка"]["status"] == "совпадает"
    assert Decimal(by_entity["ООО Ромашка"]["diff"]) == Decimal("0")

    assert by_entity["ИП Иванов"]["status"] == "расхождение"
    assert Decimal(by_entity["ИП Иванов"]["diff"]) == Decimal("500")

    assert by_entity["ЗАО Вектор"]["status"] == "только в A"
    assert by_entity["ЗАО Вектор"]["sum_b"] is None

    assert by_entity["ООО Заря"]["status"] == "только в B"
    assert by_entity["ООО Заря"]["sum_a"] is None

    assert body["summary"]["matched"] == 1
    assert body["summary"]["mismatched"] == 1
    assert body["summary"]["only_a"] == 1
    assert body["summary"]["only_b"] == 1


def test_compare_export_downloads_xlsx_with_same_rows(client):
    a = _upload_for_compare(client, _compare_file_a_bytes(), "A.xlsx")
    b = _upload_for_compare(client, _compare_file_b_bytes(), "B.xlsx")

    body = {
        "token_a": a["token"],
        "header_row_index_a": 0,
        "entity_column_a": 0,
        "metric_column_a": 1,
        "token_b": b["token"],
        "header_row_index_b": 0,
        "entity_column_b": 0,
        "metric_column_b": 1,
    }

    resp = client.post("/api/tools/compare/export", json=body)
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = openpyxl.load_workbook(BytesIO(resp.content))
    ws = wb["Расхождения"]
    header = [c.value for c in ws[1]]
    assert header == ["Сущность", "Сумма A", "Сумма B", "Разница", "Статус"]

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 4
    statuses = {r[4] for r in rows}
    assert statuses == {"совпадает", "расхождение", "только в A", "только в B"}


def test_compare_unknown_token_is_404(client):
    resp = client.post(
        "/api/tools/compare",
        json={
            "token_a": "missing",
            "header_row_index_a": 0,
            "entity_column_a": 0,
            "metric_column_a": 1,
            "token_b": "also-missing",
            "header_row_index_b": 0,
            "entity_column_b": 0,
            "metric_column_b": 1,
        },
    )
    assert resp.status_code == 404
