"""End-to-end HTTP proof of "Несколько сумм одновременно": a session can
pick MORE THAN ONE numeric column as a "sum" (e.g. "Стоимость" in rubles
AND "Вес" in kg) and every chosen sum is reconciled/broken-down/exported
independently, shown side by side rather than replacing each other. All
data below is invented test data, no real business figures.
"""

from decimal import Decimal
from io import BytesIO
from pathlib import Path

import fakeredis
import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.main import app, get_redis


@pytest.fixture
def client():
    fake_redis = fakeredis.FakeStrictRedis()
    app.dependency_overrides[get_redis] = lambda: fake_redis
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


# Менеджер, Клиент, Товар, Стоимость, Вес
_ROWS = [
    ["Менеджер", "Клиент", "Товар", "Стоимость", "Вес"],
    ["Смирнова А.", "ООО Ромашка", "носки чёрные", "4500", "12.5"],
    ["Смирнова А.", "ИП Кузнецов", "ботинки зимние", "32000", "40.0"],
    ["Кузьмин Д.", "ООО Вектор", "шапка вязаная", "8900", "6.5"],
    ["Кузьмин Д.", "ООО Вектор", "трусы муж.", "3400", "2.0"],
]


def _xlsx_bytes(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _upload(client) -> str:
    resp = client.post(
        "/api/upload",
        files={
            "file": (
                "Продажи.xlsx",
                _xlsx_bytes(_ROWS),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert resp.status_code == 200, resp.text
    return resp.json()["session_id"]


def test_columns_endpoint_accepts_and_returns_multiple_metric_columns(client):
    session_id = _upload(client)
    resp = client.post(
        f"/api/session/{session_id}/columns",
        json={
            "header_row_index": 0,
            "categorized_columns": [2],  # Товар
            "metric_columns": [3, 4],  # Стоимость, Вес
            "dimension_columns": [0],  # Менеджер
        },
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["metric_columns"] == ["Стоимость", "Вес"]

    get_resp = client.get(f"/api/session/{session_id}/columns")
    assert get_resp.json()["metric_columns"] == ["Стоимость", "Вес"]


def test_zero_metric_columns_is_a_400(client):
    session_id = _upload(client)
    resp = client.post(
        f"/api/session/{session_id}/columns",
        json={"header_row_index": 0, "categorized_columns": [2], "metric_columns": []},
    )
    assert resp.status_code == 400


def test_summary_reconciles_every_metric_independently_side_by_side(client):
    session_id = _upload(client)
    client.post(
        f"/api/session/{session_id}/columns",
        json={
            "header_row_index": 0,
            "categorized_columns": [2],
            "metric_columns": [3, 4],
            "dimension_columns": [0],
        },
    )
    client.post(f"/api/session/{session_id}/groups/Товар", json={"group_id": "socks", "name": "Носки"})
    assert client.post(
        f"/api/session/{session_id}/assign/Товар",
        json={"entity_ids": ["носки чёрные"], "group_id": "socks"},
    ).status_code == 200
    # everything else deliberately left unassigned

    resp = client.get(f"/api/session/{session_id}/summary/Товар")
    assert resp.status_code == 200
    body = resp.json()
    assert body["metric_columns"] == ["Стоимость", "Вес"]
    assert set(body["metrics"]) == {"Стоимость", "Вес"}

    cost = body["metrics"]["Стоимость"]
    weight = body["metrics"]["Вес"]

    cost_total = Decimal("4500") + Decimal("32000") + Decimal("8900") + Decimal("3400")
    weight_total = Decimal("12.5") + Decimal("40.0") + Decimal("6.5") + Decimal("2.0")

    assert Decimal(cost["rollup_totals"]["socks"]) == Decimal("4500")
    assert Decimal(cost["unassigned_total"]) == cost_total - Decimal("4500")
    assert Decimal(cost["grand_total"]) == cost_total

    # SAME tree, SAME assignment, completely independent second sum —
    # must reconcile to ITS OWN total, not be conflated with cost.
    assert Decimal(weight["rollup_totals"]["socks"]) == Decimal("12.5")
    assert Decimal(weight["unassigned_total"]) == weight_total - Decimal("12.5")
    assert Decimal(weight["grand_total"]) == weight_total


def test_breakdown_merges_every_metric_into_one_table(client):
    session_id = _upload(client)
    client.post(
        f"/api/session/{session_id}/columns",
        json={
            "header_row_index": 0,
            "categorized_columns": [2],
            "metric_columns": [3, 4],
            "dimension_columns": [0],
        },
    )
    client.post(f"/api/session/{session_id}/groups/Товар", json={"group_id": "socks", "name": "Носки"})
    client.post(
        f"/api/session/{session_id}/assign/Товар",
        json={"entity_ids": ["носки чёрные"], "group_id": "socks"},
    )

    resp = client.get(f"/api/session/{session_id}/breakdown/Товар")
    assert resp.status_code == 200
    body = resp.json()
    assert body["metric_columns"] == ["Стоимость", "Вес"]
    rows = {(r["Менеджер"], r["Группа 1"]): r for r in body["rows"]}

    smirnova_socks = rows[("Смирнова А.", "Носки")]
    assert Decimal(smirnova_socks["Стоимость"]) == Decimal("4500")
    assert Decimal(smirnova_socks["Вес"]) == Decimal("12.5")


def test_export_names_sheets_for_the_tree_x_metric_cross_product(client):
    """1 tree, 2 metrics: the FIRST metric keeps the exact old plain sheet
    names ("Итоги"/"Разбивка") for backward compatibility, the second gets
    a "— {metric}" suffix — same convention already used for a second
    TREE, now crossed with the metric axis too."""
    session_id = _upload(client)
    client.post(
        f"/api/session/{session_id}/columns",
        json={
            "header_row_index": 0,
            "categorized_columns": [2],
            "metric_columns": [3, 4],
            "dimension_columns": [0],
        },
    )
    client.post(f"/api/session/{session_id}/groups/Товар", json={"group_id": "socks", "name": "Носки"})
    client.post(
        f"/api/session/{session_id}/assign/Товар",
        json={"entity_ids": ["носки чёрные"], "group_id": "socks"},
    )

    resp = client.get(f"/api/session/{session_id}/export")
    assert resp.status_code == 200
    workbook = openpyxl.load_workbook(BytesIO(resp.content))
    assert set(workbook.sheetnames) == {"Детализация", "Итоги", "Итоги — Вес", "Разбивка", "Разбивка — Вес"}

    itogi = {row[0].value.strip(): row[1].value for row in workbook["Итоги"].iter_rows(min_row=2)}
    itogi_weight = {row[0].value.strip(): row[1].value for row in workbook["Итоги — Вес"].iter_rows(min_row=2)}

    cost_total = Decimal("4500") + Decimal("32000") + Decimal("8900") + Decimal("3400")
    weight_total = Decimal("12.5") + Decimal("40.0") + Decimal("6.5") + Decimal("2.0")
    assert Decimal(itogi["ИТОГО"]) == cost_total
    assert Decimal(itogi_weight["ИТОГО"]) == weight_total
    assert Decimal(itogi["Носки"]) == Decimal("4500")
    assert Decimal(itogi_weight["Носки"]) == Decimal("12.5")


def test_export_two_trees_two_metrics_names_every_combination(client):
    session_id = _upload(client)
    client.post(
        f"/api/session/{session_id}/columns",
        json={
            "header_row_index": 0,
            "categorized_columns": [2, 1],  # Товар (first), Клиент (second)
            "metric_columns": [3, 4],  # Стоимость (first), Вес (second)
        },
    )
    resp = client.get(f"/api/session/{session_id}/export")
    assert resp.status_code == 200
    workbook = openpyxl.load_workbook(BytesIO(resp.content))
    # Товар+Стоимость = the base pair (plain names); every other
    # combination gets its own suffixed sheet.
    assert set(workbook.sheetnames) == {
        "Детализация", "Итоги", "Итоги — Вес", "Итоги — Клиент", "Итоги — Клиент — Вес",
    }
