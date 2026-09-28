"""Tests core/hierarchical.py against a SYNTHETIC workbook built with the
exact structure confirmed on a real 1C export (see core/hierarchical.py's
module docstring and the project brief): 4 nesting levels encoded purely
via Excel cell alignment.indent on column A (0/2/4/6 = Менеджер/Клиент/
Точка/Товар), a "Итого" grand-total row at indent 0, and level-caption +
metric-header rows in the preamble. All names/numbers below are invented —
no real business data.
"""

from decimal import Decimal

import openpyxl
import pytest
import xlsxwriter

from core.hierarchical import detect_and_unflatten

# Six invented leaf (Товар) rows across 2 managers / 3 clients / 4 points.
# (manager, client, point, item, Стоимость, Количество, Вес)
_LEAVES = [
    ("Смирнова Анна", "ООО Ромашка", "Точка на Ленина", "Хлеб бородинский", 100.50, 10, 5.5),
    ("Смирнова Анна", "ООО Ромашка", "Точка на Ленина", "Батон нарезной", 200.25, 20, 8.2),
    ("Смирнова Анна", "ООО Ромашка", "Точка на Мира", "Соль поваренная", 50.00, 5, 2.0),
    ("Смирнова Анна", "ИП Кузнецов", "Точка склад", "Сахар", 300.00, 15, 15.0),
    ("Петров Игорь", "ЗАО Вектор", "Точка офис", "Мука пшеничная", 150.00, 12, 12.0),
    ("Петров Игорь", "ЗАО Вектор", "Точка офис", "Крупа гречневая", 99.99, 9, 4.5),
]
# Computed by hand with Decimal (not float sum) — see core/hierarchical.py's
# docstring reasoning on why literal 2-decimal floats round-trip exactly.
_TOTAL_COST = 900.74
_TOTAL_QTY = 71
_TOTAL_WEIGHT = 47.2


def _build_workbook(path, total_cost=_TOTAL_COST, total_qty=_TOTAL_QTY, total_weight=_TOTAL_WEIGHT,
                     include_captions=True, include_total=True, leaves=_LEAVES):
    # Built with xlsxwriter, not openpyxl.Workbook() — verified live that
    # openpyxl's own writer emits inline strings and skips xl/sharedStrings
    # .xml entirely for a small sheet like this one, which would make
    # test_sharedstrings_xml_casing_bug_is_worked_around a no-op (nothing
    # to rename). xlsxwriter (already a project dependency, used by
    # core/export.py) always writes a real xl/sharedStrings.xml, and its
    # `indent` cell format round-trips through openpyxl's
    # cell.alignment.indent exactly like a file written by Excel/1C itself
    # — confirmed live, not assumed.
    wb = xlsxwriter.Workbook(str(path))
    ws = wb.add_worksheet()
    indent_formats = {indent: wb.add_format({"indent": indent}) for indent in (0, 2, 4, 6)}

    r = 0

    def set_a(row, text, indent):
        ws.write(row, 0, text, indent_formats[indent])

    def set_metrics(row, cost, qty, weight):
        ws.write(row, 1, cost)
        ws.write(row, 2, qty)
        ws.write(row, 3, weight)

    # Plain metric header row (column A blank).
    ws.write(r, 1, "Стоимость")
    ws.write(r, 2, "Количество")
    ws.write(r, 3, "Вес")
    r += 1

    if include_total:
        set_a(r, "Итого", 0)
        set_metrics(r, total_cost, total_qty, total_weight)
        r += 1

    if include_captions:
        set_a(r, "Торговый агент", 0)
        r += 1
        set_a(r, "Контрагент", 2)
        r += 1
        set_a(r, "Торговая точка", 4)
        r += 1
        set_a(r, "Номенклатура", 6)
        r += 1

    current_manager = current_client = current_point = None
    for manager, client, point, item, cost, qty, weight in leaves:
        if manager != current_manager:
            set_a(r, manager, 0)
            r += 1
            current_manager = manager
            current_client = None
            current_point = None
        if client != current_client:
            set_a(r, client, 2)
            r += 1
            current_client = client
            current_point = None
        if point != current_point:
            set_a(r, point, 4)
            r += 1
            current_point = point
        set_a(r, item, 6)
        set_metrics(r, cost, qty, weight)
        r += 1

    wb.close()
    return path


@pytest.fixture
def hierarchical_workbook(tmp_path):
    return _build_workbook(tmp_path / "hierarchical.xlsx")


def test_detects_and_unflattens_to_one_row_per_leaf(hierarchical_workbook):
    result = detect_and_unflatten(hierarchical_workbook)
    assert result is not None
    assert result.df.height == 1 + len(_LEAVES)  # header row + one row per leaf


def test_header_row_uses_real_extracted_names_not_generic(hierarchical_workbook):
    result = detect_and_unflatten(hierarchical_workbook)
    header = result.df.row(0)
    assert header == (
        "Торговый агент", "Контрагент", "Торговая точка", "Номенклатура",
        "Стоимость", "Количество", "Вес",
    )
    # Column names on the DataFrame itself must stay generic (column_N) —
    # same raw-grid contract as core.parsing.parse_file_raw(), so the
    # existing "click the header row" UI/promote_header_row() need zero
    # changes to consume this.
    assert result.df.columns == [f"column_{i + 1}" for i in range(7)]


def test_leaf_rows_carry_the_currently_in_effect_ancestor_values(hierarchical_workbook):
    result = detect_and_unflatten(hierarchical_workbook)
    rows = result.df.rows()[1:]
    # Row for "Батон нарезной" must show its OWN manager/client/point context.
    baton = next(r for r in rows if r[3] == "Батон нарезной")
    assert baton[:4] == ("Смирнова Анна", "ООО Ромашка", "Точка на Ленина", "Батон нарезной")
    assert baton[4:] == ("200.25", "20", "8.2")

    krupa = next(r for r in rows if r[3] == "Крупа гречневая")
    assert krupa[:4] == ("Петров Игорь", "ЗАО Вектор", "Точка офис", "Крупа гречневая")


def test_leaf_sums_reconcile_with_itogo_row_via_decimal(hierarchical_workbook):
    result = detect_and_unflatten(hierarchical_workbook)
    rows = result.df.rows()[1:]
    cost_sum = sum((Decimal(r[4]) for r in rows), Decimal(0))
    qty_sum = sum((Decimal(r[5]) for r in rows), Decimal(0))
    weight_sum = sum((Decimal(r[6]) for r in rows), Decimal(0))
    assert cost_sum == Decimal(str(_TOTAL_COST))
    assert qty_sum == Decimal(str(_TOTAL_QTY))
    assert weight_sum == Decimal(str(_TOTAL_WEIGHT))


def test_no_indent_signal_returns_none(tmp_path):
    """A perfectly ordinary flat export (no grouping at all) must NOT be
    mistaken for this shape — every column-A cell at indent 0."""
    path = tmp_path / "flat.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Товар", "Сумма"])
    ws.append(["Носки", 100])
    ws.append(["Шапка", 250])
    wb.save(path)

    assert detect_and_unflatten(path) is None


def test_no_itogo_row_returns_none_not_a_guess(tmp_path):
    """Indent signal present, but nothing to validate the leaf-sum
    against — must NOT force a guess, must fall back cleanly."""
    path = _build_workbook(tmp_path / "no_total.xlsx", include_total=False, include_captions=False)
    assert detect_and_unflatten(path) is None


def test_mismatched_total_returns_none_not_forced(tmp_path):
    """The confidence gate: if the file's own Итого doesn't match the sum
    of leaves to the penny, this is NOT confidently this shape — don't
    force it, let the user fall back to manual column mapping."""
    path = _build_workbook(tmp_path / "mismatch.xlsx", total_cost=999999.99)
    assert detect_and_unflatten(path) is None


def test_missing_captions_falls_back_to_generic_level_names(tmp_path):
    """No level-caption rows in the preamble — must still detect and
    unflatten correctly (the core arithmetic doesn't depend on captions),
    just with a plain fallback name per level, never a crash."""
    path = _build_workbook(tmp_path / "no_captions.xlsx", include_captions=False)
    result = detect_and_unflatten(path)
    assert result is not None
    header = result.df.row(0)
    assert header[:4] == ("Уровень 1", "Уровень 2", "Уровень 3", "Уровень 4")
    assert result.df.height == 1 + len(_LEAVES)


def test_single_child_chain_not_mistaken_for_a_caption_ladder(tmp_path):
    """The exact edge case that motivated requiring the WHOLE ladder
    (0/2/4/6) to have blank metrics before treating it as a caption block:
    a manager with exactly one client, who has exactly one point, who has
    exactly one item, produces consecutive rows at 0/2/4/6 too — but the
    LAST one (the leaf) genuinely carries metric data, which a real
    caption-only legend row never does. Must be treated as real data, not
    swallowed as a legend, and its ancestor chain must not go missing."""
    single_chain = [
        ("Единственный Менеджер", "Единственный Клиент", "Единственная Точка", "Единственный Товар", 42.00, 3, 1.5),
    ]
    path = _build_workbook(
        tmp_path / "single_chain.xlsx",
        total_cost=42.00, total_qty=3, total_weight=1.5,
        include_captions=False, leaves=single_chain,
    )
    result = detect_and_unflatten(path)
    assert result is not None
    assert result.df.height == 2  # header + exactly one leaf
    row = result.df.row(1)
    assert row[:4] == ("Единственный Менеджер", "Единственный Клиент", "Единственная Точка", "Единственный Товар")
    assert row[4:] == ("42", "3", "1.5")


def test_sharedstrings_xml_casing_bug_is_worked_around(tmp_path):
    """Reproduces the real bug found live: some 1C exports name the
    shared-strings zip entry `xl/SharedStrings.xml` (capital S) instead of
    the OOXML-standard lowercase. A plain openpyxl.load_workbook() raises
    KeyError on that — detect_and_unflatten() must recover via the
    documented re-zip workaround, not propagate the crash."""
    import zipfile

    good_path = tmp_path / "good.xlsx"
    _build_workbook(good_path)

    bad_path = tmp_path / "bad_casing.xlsx"
    with zipfile.ZipFile(good_path, "r") as zin, zipfile.ZipFile(bad_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            name = item.filename
            if name == "xl/sharedStrings.xml":
                name = "xl/SharedStrings.xml"
            zout.writestr(name, data)

    # Sanity: plain openpyxl really does choke on this file (proves the
    # test reproduces the real bug, not a no-op).
    with pytest.raises(KeyError):
        openpyxl.load_workbook(bad_path)

    result = detect_and_unflatten(bad_path)
    assert result is not None
    assert result.df.height == 1 + len(_LEAVES)
