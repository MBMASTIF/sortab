from decimal import Decimal

import openpyxl
import polars as pl

from core.export import (
    UNASSIGNED_LABEL,
    build_breakdown_sheet,
    build_detail_sheet,
    build_multi_tree_detail_sheet,
    build_summary_sheet,
    export_workbook,
)
from core.reconcile import Row, rollup
from core.tree import TreeStore


def _sample_tree():
    store = TreeStore()
    store.add_group("clothes", "Одежда")
    store.add_group("mens", "Мужская", parent_id="clothes")
    store.add_group("socks", "Носки", parent_id="mens")
    store.add_group("underwear", "Бельё", parent_id="mens")
    store.assign_entity("носки чёрные", "socks")
    store.assign_entity("носки белые", "socks")
    store.assign_entity("трусы", "underwear")
    return store


def test_detail_sheet_adds_group_path_columns():
    store = _sample_tree()
    original = pl.DataFrame(
        {
            "Товар": ["носки чёрные", "носки белые", "трусы", "неизвестный товар"],
            "Сумма": [500, 300, 200, 999],
        }
    )

    detail = build_detail_sheet(original, "Товар", store.assignment, store.as_group_list())

    assert "Группа 1" in detail.columns
    assert "Группа 2" in detail.columns
    assert "Группа 3" in detail.columns
    row0 = detail.filter(pl.col("Товар") == "носки чёрные").row(0, named=True)
    assert row0["Группа 1"] == "Одежда"
    assert row0["Группа 2"] == "Мужская"
    assert row0["Группа 3"] == "Носки"


def test_detail_sheet_marks_unassigned_rows_explicitly():
    store = _sample_tree()
    original = pl.DataFrame({"Товар": ["неизвестный товар"], "Сумма": [999]})
    detail = build_detail_sheet(original, "Товар", store.assignment, store.as_group_list())
    assert detail.row(0, named=True)["Группа 1"] == UNASSIGNED_LABEL


def test_summary_sheet_matches_rollup_totals():
    store = _sample_tree()
    rows = [
        Row("носки чёрные", Decimal("500")),
        Row("носки белые", Decimal("300")),
        Row("трусы", Decimal("200")),
        Row("неизвестный товар", Decimal("999")),
    ]
    result = rollup(rows, store.assignment, store.as_group_list())
    summary = build_summary_sheet(result, store.as_group_list())

    as_dict = dict(zip(summary["Группа"].str.strip_chars().to_list(), summary["Сумма"].to_list()))
    assert as_dict["Носки"] == "800"
    assert as_dict["Мужская"] == "1000"
    assert as_dict["Одежда"] == "1000"
    assert as_dict[UNASSIGNED_LABEL] == "999"
    assert as_dict["ИТОГО"] == "1999"


def test_full_pipeline_exports_a_real_readable_xlsx(tmp_path):
    """End-to-end: file data -> tree -> rollup -> export -> read back with
    openpyxl and verify the actual written bytes, not just that no
    exception was raised."""
    store = _sample_tree()
    original = pl.DataFrame(
        {
            "Товар": ["носки чёрные", "носки белые", "трусы", "неизвестный товар"],
            "Сумма": [500, 300, 200, 999],
        }
    )
    rows = [Row(t, Decimal(str(s))) for t, s in zip(original["Товар"], original["Сумма"])]
    result = rollup(rows, store.assignment, store.as_group_list())

    detail_df = build_detail_sheet(original, "Товар", store.assignment, store.as_group_list())
    summary_df = build_summary_sheet(result, store.as_group_list())

    out_path = tmp_path / "result.xlsx"
    export_workbook(str(out_path), detail_df, summary_df)

    book = openpyxl.load_workbook(out_path)
    assert book.sheetnames == ["Детализация", "Итоги"]

    detail_sheet = book["Детализация"]
    header = [c.value for c in detail_sheet[1]]
    assert "Группа 3" in header

    summary_sheet = book["Итоги"]
    summary_values = [(row[0].value, row[1].value) for row in summary_sheet.iter_rows(min_row=2)]
    summary_dict = {k.strip(): v for k, v in summary_values}
    assert summary_dict["ИТОГО"] == "1999"
    assert summary_dict[UNASSIGNED_LABEL] == "999"


# ---------------------------------------------------------------------
# Multiple independent trees in one session (see core/session.py) — the
# detail sheet grows a path per tree, and a workbook can carry more than
# one "Итоги"/"Разбивка" sheet, one pair per categorized column.
# ---------------------------------------------------------------------


def _client_tree():
    store = TreeStore()
    store.add_group("big", "Крупные")
    store.add_group("small", "Мелкие")
    store.assign_entity("ООО Ромашка", "big")
    store.assign_entity("ИП Сидоров", "small")
    return store


def test_multi_tree_detail_sheet_adds_a_path_per_tree_prefixed_by_column():
    product_tree = _sample_tree()  # Одежда > Мужская > {Носки, Бельё}
    client_tree = _client_tree()
    original = pl.DataFrame(
        {
            "Менеджер": ["Иванов", "Петров", "Иванов", "Петров"],
            "Товар": ["носки чёрные", "носки белые", "трусы", "неизвестный товар"],
            "Клиент": ["ООО Ромашка", "ИП Сидоров", "ООО Ромашка", "ООО Вектор"],
            "Сумма": [500, 300, 200, 999],
        }
    )

    detail = build_multi_tree_detail_sheet(
        original,
        ["Товар", "Клиент"],
        {"Товар": product_tree, "Клиент": client_tree},
        dimension_columns=["Менеджер"],
    )

    assert detail.columns[0] == "Менеджер"  # dimension column pulled to front
    assert "Товар: Группа 1" in detail.columns
    assert "Товар: Группа 3" in detail.columns
    assert "Клиент: Группа 1" in detail.columns
    assert detail.height == original.height  # no row duplication/loss from the per-tree joins

    row0 = detail.filter(pl.col("Товар") == "носки чёрные").row(0, named=True)
    assert row0["Товар: Группа 1"] == "Одежда"
    assert row0["Товар: Группа 3"] == "Носки"
    assert row0["Клиент: Группа 1"] == "Крупные"

    row_unassigned = detail.filter(pl.col("Товар") == "неизвестный товар").row(0, named=True)
    assert row_unassigned["Товар: Группа 1"] == UNASSIGNED_LABEL
    assert row_unassigned["Клиент: Группа 1"] == UNASSIGNED_LABEL


def test_breakdown_sheet_sums_metric_per_dimension_and_group_path_natively():
    """Expected sums hand-computed independently, not trusted from the
    implementation's own arithmetic."""
    store = _sample_tree()  # Одежда > Мужская > {Носки, Бельё}
    original = pl.DataFrame(
        {
            "Товар": ["носки чёрные", "носки чёрные", "трусы", "шапка"],
            "Менеджер": ["Иванов", "Петров", "Иванов", "Петров"],
            "Сумма": [500, 300, 200, 999],
        }
    )
    # носки чёрные -> socks, трусы -> underwear, шапка left unassigned (per _sample_tree)

    breakdown = build_breakdown_sheet(
        original, "Товар", "Сумма", ["Менеджер"], store.assignment, store.as_group_list()
    )

    as_map = {
        (r["Менеджер"], r["Группа 1"], r["Группа 2"], r["Группа 3"]): r["Сумма"]
        for r in breakdown.iter_rows(named=True)
    }
    assert as_map[("Иванов", "Одежда", "Мужская", "Носки")] == "500"
    assert as_map[("Петров", "Одежда", "Мужская", "Носки")] == "300"
    assert as_map[("Иванов", "Одежда", "Мужская", "Бельё")] == "200"
    assert as_map[("Петров", UNASSIGNED_LABEL, "", "")] == "999"
    assert breakdown.height == 4  # 4 distinct (Менеджер, path) combos, none merged incorrectly


def test_export_workbook_writes_extra_summary_and_breakdown_sheets(tmp_path):
    store = _sample_tree()
    original = pl.DataFrame({"Товар": ["носки чёрные", "трусы"], "Сумма": [500, 200]})
    rows = [Row(t, Decimal(str(s))) for t, s in zip(original["Товар"], original["Сумма"])]
    result = rollup(rows, store.assignment, store.as_group_list())

    detail_df = build_detail_sheet(original, "Товар", store.assignment, store.as_group_list())
    summary_df = build_summary_sheet(result, store.as_group_list())
    breakdown_df = build_breakdown_sheet(original, "Товар", "Сумма", [], store.assignment, store.as_group_list())

    second_store = TreeStore()
    second_store.add_group("all", "Всё")
    second_store.assign_entity("носки чёрные", "all")
    second_store.assign_entity("трусы", "all")
    second_result = rollup(rows, second_store.assignment, second_store.as_group_list())
    second_summary_df = build_summary_sheet(second_result, second_store.as_group_list())

    out_path = tmp_path / "multi.xlsx"
    export_workbook(
        str(out_path),
        detail_df,
        summary_df,
        breakdown_df,
        extra_summary_sheets={"Итоги — Клиент": second_summary_df},
    )

    book = openpyxl.load_workbook(out_path)
    assert set(book.sheetnames) == {"Детализация", "Итоги", "Разбивка", "Итоги — Клиент"}
